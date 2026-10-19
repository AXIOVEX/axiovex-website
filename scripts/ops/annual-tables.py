"""Validate/rebuild existing MCDA, IPEDS and Census PSEO views from primary files.
No inferred suppressed cells, cohorts, populations, or projection horizons.
"""
import csv
import gzip
import io
import json
from pathlib import Path
import re
import sys
import zipfile
import openpyxl
from datetime import date
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / 'ops-output/annual-sources'
def load(path):
    return json.loads((ROOT / path).read_text(encoding='utf-8'))
def save(path, obj):
    # No retrieval-date-only rewrites.
    before = load(path)
    if {k:v for k,v in before.items() if k != 'retrievedOn'} != {k:v for k,v in obj.items() if k != 'retrievedOn'}:
        if 'retrievedOn' in obj: obj['retrievedOn']=date.today().isoformat()
        if path=='data/outlook.json':
            for key in ['blsProjections','michiganProjections','michiganOccupations']:
                if before.get(key)!=obj.get(key): obj[key]['retrievedOn']=date.today().isoformat()
            for f in obj['education']['figures']:
                if next((v for v in before['education']['figures'] if v['id']==f['id']),None)!=f: f['retrievedOn']=date.today().isoformat()
        (ROOT / path).write_bytes((json.dumps(obj, indent=2, ensure_ascii=False)+'\n').encode())
def csv_file(path):
    raw = path.read_bytes()
    if path.suffix == '.gz': raw = gzip.decompress(raw)
    if path.suffix == '.zip':
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            name = next(n for n in z.namelist() if n.lower().endswith('.csv'))
            raw = z.read(name)
    return list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
def number(value):
    v = float(str(value).replace(',', ''))
    if not (abs(v) < 1e12): raise ValueError('Nonfinite or implausible number')
    return int(v) if v.is_integer() else v

def mcda():
    outlook = load('data/outlook.json')
    for index, source in enumerate(outlook['michiganOccupations']['sourceFiles']):
        wb = openpyxl.load_workbook(INPUT/f'mcda-{index}.xlsx', read_only=True, data_only=True)
        rows = list(wb.worksheets[0].values)
        header = next(i for i,r in enumerate(rows) if str(r[0]).strip() == 'SOC Code')
        text = ' '.join(str(v) for r in rows[:header] for v in r if v)
        period = re.search(r'(20\d\d)\s*(?:-|–|to)\s*(20\d\d)', text)
        if not period: raise ValueError('MCDA projection horizon absent')
        start,end = period.groups()
        if int(end)-int(start) != 10: raise ValueError('MCDA horizon not ten years')
        lookup = {str(r[0]).strip():r for r in rows[header+1:] if re.fullmatch(r'\d{2}-\d{4}',str(r[0]).strip())}
        group = next(g for g in outlook['michiganOccupations']['groups'] if g['geography'] == source['geography'])
        for old in group['rows']:
            r = lookup[old['soc']]
            base, projected, pct, openings = number(r[2]), number(r[3]), number(r[5]), number(r[10])
            if base <= 0 or projected < 0 or openings < 0: raise ValueError('Invalid MCDA levels')
            # MCDA publishes change fractions, not percentage points.
            if abs(pct*100 - (projected/base-1)*100) > .6: raise ValueError('MCDA percent/level reconciliation failed')
            old.update(occupation=str(r[1]).strip(), baseEmployment=base, projectedEmployment=projected, pctChange=round(pct*100,1), annualOpenings=openings)
        group['period'] = f'{start}-{end}'
        group['horizon'] = f'{start}–{end[-2:]}'
        source['period'] = group['period']
    outlook['michiganOccupations']['viaStudy'] = False
    save('data/outlook.json',outlook)

def bls():
    outlook=load('data/outlook.json')
    wb=openpyxl.load_workbook(INPUT/'bls-occupations.xlsx',read_only=True,data_only=True)
    horizons=[]
    for name,key in [('Table 1.3','rise'),('Table 1.5','fall')]:
        rows=list(wb[name].values)
        years=re.search(r'(20\d\d).*?(20\d\d)',str(rows[0][0]))
        if not years or int(years[2])-int(years[1])!=10: raise ValueError('BLS workbook horizon changed')
        horizons.append(f'{years[1]}–{years[2][-2:]}')
        selected=[r for r in rows[2:] if re.fullmatch(r'\d{2}-\d{4}',str(r[1])) and r[1]!='00-0000'][:5]
        if len(selected)!=5: raise ValueError('BLS workbook table coverage changed')
        outlook['blsProjections'][key]=[{'name':re.sub(r'\(\d+\)','',str(r[0])).strip(),'jobsChange':round(number(r[4])*1000),'pctChange':number(r[5])} for r in selected]
    if horizons[0]!=horizons[1]: raise ValueError('BLS workbook horizons disagree')
    outlook['blsProjections'].update(horizon=horizons[0],vintage=horizons[0],sourceFileUrl='https://www.bls.gov/emp/ind-occ-matrix/occupation.xlsx')
    save('data/outlook.json',outlook)

def bls_bulk():
    def tsv(name):
        with (INPUT/name).open() as f:
            return [{k.strip():v.strip() for k,v in r.items()} for r in csv.DictReader(f,delimiter='\t')]
    dates=next(r for r in tsv('ep.dates') if r['data_source']=='optd')
    base,end=dates['base_year'],dates['proj_year']
    if int(end)-int(base)!=10: raise ValueError('BLS bulk projection horizon changed')
    if not any(r['ind_code']=='TE1000' and r['ind_title']=='Total, all industries' for r in tsv('ep.industry')): raise ValueError('BLS national industry scope changed')
    series={r['series_id']:r for r in tsv('ep.series') if r['ind_code']=='TE1000' and r['occ_type']=='L'}
    names={r['occ_code']:r['occ_title'] for r in tsv('ep.occupation')}
    levels={r['series_id']:number(r['value']) for r in tsv('ep.data.1.AllData') if r['series_id'] in series and r['year']==base and r['period']=='A01'}
    aspects={}
    for r in tsv('ep.aspect'):
        if r['series_id'] in series and r['year']==base and r['period']=='A01': aspects.setdefault(r['series_id'],{})[r['aspect_type']]=number(r['value'])
    # Validate bulk aspect-code meaning against the API's named measures.
    api=json.loads((INPUT/'ep-api-validation.json').read_text())
    observation=api['Results']['series'][0]['data'][0]
    named={r['name']:number(r['value']) for r in observation['aspects']}
    sample=aspects['EPU291171TE1000']
    if observation['year']!=base or sample['A1']!=named['Numerical employment change'] or sample['A2']!=named['Percent employment change'] or sample['PR']!=named['Projected employment']: raise ValueError('BLS named API and bulk measures disagree')
    rows=[]
    for sid,s in series.items():
        a=aspects[sid];level=levels[sid]
        if abs((a['PR']-level)-a['A1'])>.21 or (level>1 and abs(a['A2']-a['A1']/level*100)>(.1+10/level)): raise ValueError('BLS bulk levels/change reconciliation failed')
        rows.append({'name':names[s['occ_code']],'jobsChange':round(a['A1']*1000),'pctChange':a['A2']})
    if len(rows)<500: raise ValueError('BLS detailed occupation coverage incomplete')
    outlook=load('data/outlook.json');group=outlook['blsProjections']
    group.update(rise=sorted(rows,key=lambda r:-r['pctChange'])[:5],fall=sorted(rows,key=lambda r:r['pctChange'])[:5],horizon=f'{base}–{end[-2:]}',vintage=f'{base}–{end[-2:]}',sourceFileUrl='https://download.bls.gov/pub/time.series/ep/')
    save('data/outlook.json',outlook)

def industry():
    outlook=load('data/outlook.json')
    rows=list(openpyxl.load_workbook(INPUT/'mcda-industry.xlsx',read_only=True,data_only=True).worksheets[0].values)
    period=re.search(r'(20\d\d)-(20\d\d)',str(rows[0][0]))
    if not period or int(period[2])-int(period[1])!=10: raise ValueError('MCDA industry horizon missing')
    # MCDA publishes mixed three/four-digit detail; exclude totals and broad aggregates.
    selected=[]
    for r in rows[3:]:
        code=str(r[0])
        if not re.fullmatch(r'\d{6}',code) or code.endswith('0000') or code.startswith(('00','10')) or not isinstance(r[5],(float,int)): continue
        base,projected,change,pct=map(number,[r[2],r[3],r[4],r[5]])
        if base<=0 or abs(projected-base-change)>10 or abs((projected/base-1)-pct)>.006: raise ValueError('MCDA industry reconciliation failed')
        selected.append({'name':str(r[1]).strip(),'jobsChange':change,'pctChange':round(pct*100,1)})
    if len(selected)<50: raise ValueError('MCDA industry coverage missing')
    group=outlook['michiganProjections']
    group.update(rise=sorted(selected,key=lambda r:-r['pctChange'])[:5],fall=sorted(selected,key=lambda r:r['pctChange'])[:5],horizon=f'{period[1]}–{period[2]}',vintage=f'{period[1]}–{period[2]}')
    save('data/outlook.json',outlook)

def pupils():
    outlook=load('data/outlook.json')
    source=json.loads((INPUT/'pupil-source.json').read_text())
    content='\n'.join(p.extract_text() for p in PdfReader(INPUT/'pupils.pdf').pages)
    current=date.today().year if date.today().month>=10 else date.today().year-1
    fiscal=f'FY {current}-{str(current+1)[-2:]}'
    block=re.search(re.escape(fiscal)+r'\s+LOCAL DISTRICTS[\s\S]*?TOTAL\s+([\d,]+)\s+\(?[\d,]+\)?\s+([\d,]+)',content)
    if not block or not re.search(r'PROPOSED.*?CONSENSUS',content,re.S): raise ValueError('Pupil handout schema/estimate label changed')
    value=number(block[2])
    if not 1000000<value<2000000: raise ValueError('Pupil estimate out of range')
    figure=next(f for f in outlook['education']['figures'] if f['id']=='pupil-membership')
    figure.update(value=value,valueText=f'{value:,}',period=fiscal.replace(' ','').replace('-','–'),publisher='Michigan consensus revenue conference',publisherTag=source['text'].upper(),sourceUrl=source['url'],noteText='proposed conference estimate, not an audited count',estimate=True)
    save('data/outlook.json',outlook)

def ipeds(year):
    before = load('data/geo/ipeds-institutions.json')
    schools = [r for r in csv_file(INPUT/'ipeds-hd.zip') if r.get('STABBR') == 'MI']
    if len(schools)<100: raise ValueError('IPEDS Michigan coverage unexpectedly small')
    by_id = {r['UNITID']:{'unitid':int(r['UNITID']), 'name':r['INSTNM'], 'city':r['CITY'], 'lat':number(r['LATITUDE']) if r.get('LATITUDE','').strip() else None, 'lon':number(r['LONGITUD']) if r.get('LONGITUD','').strip() else None, 'totalCompletions':0,'byFamily':{}} for r in schools}
    for r in csv_file(INPUT/'ipeds-completions.zip'):
        if r['UNITID'] not in by_id or r.get('MAJORNUM') != '1' or not re.fullmatch(r'\d{2}\.\d{4}',r['CIPCODE']) or r['CIPCODE'].startswith('99.') or r.get('AWLEVEL') == '99': continue
        n = number(r['CTOTALT'])
        if n < 0: raise ValueError('IPEDS negative completions')
        school = by_id[r['UNITID']];family=r['CIPCODE'][:2]
        school['totalCompletions'] += n
        school['byFamily'][family] = school['byFamily'].get(family,0)+n
    if sum(s['totalCompletions'] for s in by_id.values()) < 10000: raise ValueError('IPEDS incomplete completions coverage')
    before['institutions'] = sorted(by_id.values(),key=lambda s:s['name'].casefold())
    before['sourceUrl'] = f'https://nces.ed.gov/ipeds/datacenter/data/HD{year}.zip · https://nces.ed.gov/ipeds/datacenter/data/C{year}_A.zip'
    before['vintage'] = f'{year} survey cycle; completions are {year-1}-{str(year)[-2:]} awards'
    before['source'] = f'NCES IPEDS Data Center bulk files: HD{year} + C{year}_A (Completions)'
    before['grain']=f'One entry per Michigan institution (STABBR = MI in HD{year}); published coordinates; first-major awards summed over six-digit CIP detail and award levels. CIP 99 and award-level 99 totals are excluded. Second-major awards are excluded to avoid counting a degree twice.'
    save('data/geo/ipeds-institutions.json',before)

def pseo():
    before = load('data/geo/pseo-pipelines.json')
    earnings=csv_file(INPUT/'pseoe_mi.csv.gz'); flows=csv_file(INPUT/'pseof_mi.csv.gz')
    def selected(r): return r['degree_level']=='05' and r['grad_cohort']=='0000' and r['geo_level']=='N' and r['geography']=='00'
    def one(rows,**criteria):
        hits=[r for r in rows if selected(r) and all(r.get(k)==v for k,v in criteria.items())]
        if len(hits)!=1: raise ValueError('PSEO ambiguous or missing published grain')
        return hits[0]
    def published(r,field,status):
        if r.get(status)!='1': raise ValueError('PSEO selected value is not disclosed')
        n=number(r[field])
        if n<0: raise ValueError('PSEO negative measure')
        return n
    # Preserve the explicit displayed selection; recalculate and rank that view.
    for old in before['statewideFlows']['rows']:
        f=one(flows,inst_level='S',institution='26',cip_level='2',cipcode=old['cip'],ind_level='S',industry=old['industry'])
        e=one(earnings,inst_level='S',institution='26',cip_level='2',cipcode=old['cip'],ind_level='A',industry='00')
        old['y1GradsEmp']=published(f,'y1_grads_emp','status_y1_grads_emp')
        old['y1GradsEmpInstate']=published(f,'y1_grads_emp_instate','status_y1_grads_emp_instate')
        old['y1MedianEarningsField']=published(e,'y1_p50_earnings','status_y1_earnings')
    before['statewideFlows']['label']=before['statewideFlows']['label'].replace('top field-to-industry flows','selected field-to-industry flows')
    before['statewideFlows']['rows'].sort(key=lambda r:-r['y1GradsEmp'])
    spotlight=before['umichSpotlight']
    e=one(earnings,inst_level='I',institution=spotlight['opeid'],cip_level='A',ind_level='A',industry='00')
    f=one(flows,inst_level='I',institution=spotlight['opeid'],cip_level='A',ind_level='A',industry='00')
    for year in [1,5,10]:
        k=f'y{year}'; spotlight['earnings'][k]={p:published(e,f'{k}_{p}_earnings',f'status_{k}_earnings') for p in ['p25','p50','p75']}
        spotlight['earnings'][k]['gradsEarn']=published(e,f'{k}_grads_earn',f'status_{k}_grads_earn')
        emp=published(f,f'{k}_grads_emp',f'status_{k}_grads_emp'); ins=published(f,f'{k}_grads_emp_instate',f'status_{k}_grads_emp_instate')
        if not emp or ins>emp: raise ValueError('PSEO invalid employment denominator')
        spotlight['employment'][k]={'emp':emp,'instate':ins,'sharePct':round(ins/emp*100,1)}
    for old in spotlight['topIndustryFlows']:
        r=one(flows,inst_level='I',institution=spotlight['opeid'],cip_level='A',ind_level='S',industry=old['industry'])
        old['y1GradsEmp']=published(r,'y1_grads_emp','status_y1_grads_emp');old['y1GradsEmpInstate']=published(r,'y1_grads_emp_instate','status_y1_grads_emp_instate')
    version=(INPUT/'pseo-version.txt').read_text().splitlines()[0]
    m=re.search(r'(\d{4}-\d{4}).*?(20\d\dQ[1-4])',version)
    if not m: raise ValueError('PSEO release/cohort version absent')
    before['vintage']=f'Release R{m[2]} ({version}); graduate cohorts {m[1]} pooled'
    partners=(INPUT/'pseo-partners.txt').read_text()
    if 'University of Michigan' not in partners or '10% of statewide graduates covered (2015 estimate)' not in partners: raise ValueError('PSEO coverage changed; update the scope note before publication')
    before['coverage']['coverageNote']='Census lists University of Michigan as the Michigan data contributor; coverage is 10% of statewide graduates in its 2015 estimate, not a current statewide coverage measure. State aggregate rows are published by Census, not computed by summing institutions.'
    save('data/geo/pseo-pipelines.json',before)

if __name__=='__main__':
    mode=sys.argv[1]
    if mode=='mcda': mcda()
    elif mode=='bls': bls()
    elif mode=='bls-bulk': bls_bulk()
    elif mode=='industry': industry()
    elif mode=='pupils': pupils()
    elif mode=='ipeds': ipeds(int(sys.argv[2]))
    elif mode=='pseo': pseo()
    else: raise ValueError('Unknown annual source parser')
